# Week 3 教學簡報：套用 Week 2 簡報的版面，只換文字。
# 用法：python3 scripts/build-week3-slides.py
# 每頁的文字格數必須與 Week 2 同一頁一致（docs/WEEK3_SLIDES_SOURCE.json 依此編寫），不一致會直接報錯，不會產生半套的簡報。
import copy, json, re, sys
from pptx import Presentation

SRC = 'materials/Week2_教學簡報.pptx'
OUT = 'materials/Week3_教學簡報.pptx'
HEADER_OLD = 'WEEK 2  /  問題探索與選題'
HEADER_NEW = 'WEEK 3  /  使用者訪談'

slides = json.load(open('docs/WEEK3_SLIDES_SOURCE.json', encoding='utf-8'))
prs = Presentation(SRC)
if len(slides) != len(prs.slides):
    sys.exit(f'頁數不符：JSON {len(slides)}，範本 {len(prs.slides)}')

def set_text(shape, text):
    tf = shape.text_frame
    lines = text.split('\n')
    paras = list(tf.paragraphs)
    first = paras[0]
    # 多的段落刪掉
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    template = copy.deepcopy(first._p)
    def fill(p, line):
        runs = p.runs
        if not runs:
            p.text = line; return
        runs[0].text = line
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    fill(first, lines[0])
    for line in lines[1:]:
        new_p = copy.deepcopy(template)
        tf._txBody.append(new_p)
        p = tf.paragraphs[-1]
        fill(p, line)

for idx, (slide, texts) in enumerate(zip(prs.slides, slides), 1):
    body = []
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if not t.strip(): continue
        if t == HEADER_OLD: set_text(sh, HEADER_NEW); continue
        if t == 'Startup Journey Lab': continue
        if re.fullmatch(r'\d{2} / \d{2}', t): continue
        body.append(sh)
    if len(body) != len(texts):
        sys.exit(f'第 {idx} 頁格數不符：範本 {len(body)} 格，JSON {len(texts)} 格')
    for sh, text in zip(body, texts):
        set_text(sh, text)

prs.save(OUT)
print('寫出', OUT, len(prs.slides), '頁')
