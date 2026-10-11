# 用 Week 2 簡報當版面範本，依「頁面計畫」產生任何一週的簡報：只換文字，可重複使用某一頁的版面（例如多幾頁 YC 重點）。
# 用法：python3 scripts/build-slides.py --source docs/WEEK3_SLIDES_SOURCE.json --out materials/Week3_教學簡報.pptx
# JSON 兩種寫法：
#   1) 舊式：25 個陣列，位置對應範本第 1–25 頁。
#   2) 新式：{"header":"WEEK 3  /  使用者訪談","slides":[{"template":1,"texts":[...]}, ...]}，template 是範本頁碼（1 起算），可重複。
# 每頁的文字格數必須與範本該頁一致，不一致直接報錯。
import argparse, copy, json, re, sys
from pptx import Presentation

ap = argparse.ArgumentParser()
ap.add_argument('--template', default='materials/_slide_template.pptx')
ap.add_argument('--source', required=True)
ap.add_argument('--out', required=True)
ap.add_argument('--header', default=None, help='覆寫頁首文字（舊式 JSON 需要）')
args = ap.parse_args()

HEADER_OLD = 'WEEK 2  /  問題探索與選題'
raw = json.load(open(args.source, encoding='utf-8'))
if isinstance(raw, list):
    plan = [{'template': i + 1, 'texts': t} for i, t in enumerate(raw)]
    header = args.header or HEADER_OLD
else:
    plan = raw['slides']; header = args.header or raw.get('header') or HEADER_OLD

prs = Presentation(args.template)
tmpl = list(prs.slides)
if any(p['template'] < 1 or p['template'] > len(tmpl) for p in plan):
    sys.exit('template 頁碼超出範本範圍')

from pptx.oxml.ns import qn
def clone(src):
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        new.shapes._spTree.append(copy.deepcopy(shp._element))
    # 深色頁的背景是 <p:bg>，不是 shape，要另外複製
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None:
        old = new._element.cSld.find(qn('p:bg'))
        if old is not None: new._element.cSld.remove(old)
        new._element.cSld.insert(0, copy.deepcopy(bg))
    return new

# 先依計畫複製，再把原本 25 頁移除
new_slides = [clone(tmpl[p['template'] - 1]) for p in plan]
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[:len(tmpl)]:
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)

def set_text(shape, text):
    tf = shape.text_frame; lines = text.split('\n'); paras = list(tf.paragraphs); first = paras[0]
    for p in paras[1:]: p._p.getparent().remove(p._p)
    template = copy.deepcopy(first._p)
    def fill(p, line):
        runs = p.runs
        if not runs: p.text = line; return
        runs[0].text = line
        for r in runs[1:]: r._r.getparent().remove(r._r)
    fill(first, lines[0])
    for line in lines[1:]:
        tf._txBody.append(copy.deepcopy(template)); fill(tf.paragraphs[-1], line)

total = len(new_slides)
for idx, (slide, p) in enumerate(zip(new_slides, plan), 1):
    body = []
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if not t.strip(): continue
        if t == HEADER_OLD or t == header: set_text(sh, header); continue
        if t == 'Startup Journey Lab': continue
        if re.fullmatch(r'\d{2} / \d{2}', t): set_text(sh, f'{idx:02d} / {total}'); continue
        body.append(sh)
    if len(body) != len(p['texts']):
        sys.exit(f'第 {idx} 頁（範本第 {p["template"]} 頁）格數不符：範本 {len(body)} 格，JSON {len(p["texts"])} 格')
    for sh, text in zip(body, p['texts']): set_text(sh, text)

prs.save(args.out)
print('寫出', args.out, total, '頁')
